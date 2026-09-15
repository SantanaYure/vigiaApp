from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\OneDrive\Área de Trabalho\vigiaApp")
OUT = ROOT / "output" / "InsurMinds_Desafio5_Relatorio_VigiaApp.docx"
LOGO = Path(r"C:\Users\Public\Documents\vigia-report\logo-000.jpg")

NAVY = "173650"
BLUE = "2878B8"
TEAL = "08A9A6"
GRAY = "66727D"
LIGHT = "EEF4F7"
PALE_BLUE = "E8F2F9"
WHITE = "FFFFFF"
BLACK = "202428"
GREEN = "1D7A55"
AMBER = "A7660B"
RED = "B43C3C"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                edge.set(qn(f"w:{key}"), str(edge_data[key]))


def keep_table_rows(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=9.5, bold=False, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, size=9.5, bold=False, color=BLACK, italic=False, font="Arial"):
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, color=color, font=font)
    run.italic = italic
    return run


def add_hyperlink(paragraph, text: str, url: str, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(r_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def format_paragraph(paragraph, before=0, after=5, line=1.12, align=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if keep:
        fmt.keep_with_next = True
    return paragraph


def add_body(doc, text: str, *, bold_prefix: str | None = None, after=5, italic=False):
    p = doc.add_paragraph()
    format_paragraph(p, after=after, line=1.13)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True)
        add_text(p, text[len(bold_prefix):], italic=italic)
    else:
        add_text(p, text, italic=italic)
    return p


def add_bullet(doc, text: str, *, level=0, color=BLACK, after=2.5):
    p = doc.add_paragraph()
    format_paragraph(p, after=after, line=1.08)
    p.paragraph_format.left_indent = Cm(0.45 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    add_text(p, "• ", color=BLUE, bold=True)
    add_text(p, text, color=color)
    return p


def add_numbered(doc, number: str, title: str, text: str):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(1.05)
    table.columns[1].width = Cm(15.6)
    left, right = table.rows[0].cells
    set_cell_shading(left, BLUE)
    set_cell_shading(right, PALE_BLUE)
    for cell in (left, right):
        set_cell_margins(cell, top=90, start=120, bottom=90, end=120)
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, number, size=12, bold=True, color=WHITE)
    p = right.paragraphs[0]
    format_paragraph(p, after=0, line=1.0)
    add_text(p, title + " — ", size=9.5, bold=True, color=NAVY)
    add_text(p, text, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text: str, level=2, before=7, after=4):
    p = doc.add_paragraph()
    format_paragraph(p, before=before, after=after, line=1.0, keep=True)
    size = 12.2 if level == 2 else 10.5
    add_text(p, text, size=size, bold=True, color=BLACK)
    return p


def add_page_title(doc, kicker: str, title: str, subtitle: str | None = None, *, page_break=True, force_logo=False):
    if page_break:
        doc.add_page_break()
    if force_logo:
        p = doc.add_paragraph()
        format_paragraph(p, before=0, after=5, line=1.0, keep=True)
        p.add_run().add_picture(str(LOGO), width=Cm(2.45))
    p = doc.add_paragraph()
    format_paragraph(p, before=1, after=4, line=1.0, keep=True)
    add_text(p, kicker.upper(), size=8.3, bold=True, color=BLUE)
    p = doc.add_paragraph()
    format_paragraph(p, before=0, after=8, line=0.95, keep=True)
    add_text(p, title, size=21, bold=True, color=NAVY, font="Arial")
    if subtitle:
        p = doc.add_paragraph()
        format_paragraph(p, after=10, line=1.1)
        add_text(p, subtitle, size=10.2, color=GRAY)


def add_callout(doc, label: str, text: str, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    set_cell_border(cell, start={"val": "single", "sz": "18", "color": accent}, top={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"})
    p = cell.paragraphs[0]
    format_paragraph(p, after=0, line=1.12)
    add_text(p, label + " ", size=9.5, bold=True, color=NAVY)
    add_text(p, text, size=9.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers: Iterable[str], rows: Iterable[Iterable[str]], widths=None, font_size=8.4):
    headers = list(headers)
    rows = [list(r) for r in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Cm(width)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=90, start=95, bottom=90, end=95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        format_paragraph(p, after=0, line=1.0)
        add_text(p, header, size=font_size, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if ridx % 2 == 0 else LIGHT
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, fill)
            set_cell_margins(cell, top=78, start=95, bottom=78, end=95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            format_paragraph(p, after=0, line=1.02)
            add_text(p, str(value), size=font_size, color=BLACK)
    keep_table_rows(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_metric_cards(doc, cards):
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (value, label, color) in enumerate(cards):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        set_cell_margins(cell, top=130, start=100, bottom=125, end=100)
        set_cell_border(cell, top={"val": "single", "sz": "14", "color": color}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, after=2, line=1.0)
        add_text(p, value, size=17, bold=True, color=color)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, after=0, line=1.0)
        add_text(p, label, size=7.8, bold=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_architecture(doc):
    widths = [3.15, 0.45, 3.15, 0.45, 3.15, 0.45, 3.15, 0.45, 3.15]
    table = doc.add_table(rows=1, cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    items = [
        ("FONTES", "INMET\nIBGE", BLUE),
        ("COLETA", "Normalização\ne revisão", NAVY),
        ("DECISÃO", "Elegibilidade\ne cobertura", TEAL),
        ("GERAÇÃO", "Gemini 3.5\n→ Groq OSS", BLUE),
        ("ENTREGA", "API · UI\nSQLite", NAVY),
    ]
    item_idx = 0
    for idx, cell in enumerate(table.rows[0].cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, start=70, bottom=120, end=70)
        if idx % 2 == 1:
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, "›", size=17, bold=True, color=GRAY)
            continue
        label, text, color = items[item_idx]
        item_idx += 1
        set_cell_shading(cell, color)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, after=4, line=1.0)
        add_text(p, label, size=7.5, bold=True, color=WHITE)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, after=0, line=1.0)
        add_text(p, text, size=9.1, bold=True, color=WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_check(doc, status: str, requirement: str, evidence: str):
    color = GREEN if status == "ATENDE" else AMBER
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(2.2)
    table.columns[1].width = Cm(5.1)
    table.columns[2].width = Cm(9.2)
    cells = table.rows[0].cells
    for c in cells:
        set_cell_margins(c, top=75, start=95, bottom=75, end=95)
        set_cell_shading(c, LIGHT)
    p = cells[0].paragraphs[0]
    add_text(p, status, size=8, bold=True, color=color)
    p = cells[1].paragraphs[0]
    add_text(p, requirement, size=8.5, bold=True, color=NAVY)
    p = cells[2].paragraphs[0]
    add_text(p, evidence, size=8.3)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.72)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = False

    doc.settings.odd_and_even_pages_header_footer = True
    for header in (section.header, section.even_page_header):
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(LOGO), width=Cm(2.45))

    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "VigiaApp · InsurMinds · Desafio 5", size=8.2, color=GRAY)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    core = doc.core_properties
    core.title = "InsurMinds — Desafio 5 | Relatório Técnico do VigiaApp"
    core.subject = "Ferramenta Inteligente para Comunicação Proativa com o Segurado"
    core.author = "VIL — Visionary Insurance Lab"
    core.keywords = "InsurMinds, VigiaApp, riscos climáticos, seguros, inteligência artificial"


def configure_body_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.72)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = False
    section.footer.is_linked_to_previous = False
    for header in (section.header, section.even_page_header):
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(LOGO), width=Cm(2.45))
    footer = section.footer
    footer.paragraphs[0].clear()


def build_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(32)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=1, line=1.0)
    add_text(p, "InsurMinds", size=16, bold=True, color=NAVY, font="Arial")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=18, line=1.0)
    add_text(p, "Inteligência Artificial Aplicada a Seguros", size=10.5, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=8, line=1.0)
    add_text(p, "Desafio 5", size=13.5, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=18, line=1.0)
    add_text(p, "Ferramenta Inteligente para Comunicação\nProativa com o Segurado", size=18, bold=True, color=NAVY, font="Arial")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=2, line=1.0)
    add_text(p, "Projeto: VigiaApp", size=11.5, bold=True, color=BLACK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=25, line=1.0)
    add_text(p, "Grupo: VIL — Visionary Insurance Lab", size=10.2, color=GRAY)

    p = doc.add_paragraph()
    format_paragraph(p, after=4, line=1.0)
    add_text(p, "Representante", size=9, bold=True, color=BLUE)
    add_body(doc, "Yure Santana | (71) 99140-8574 | yure.s.santana@outlook.com", after=12)

    p = doc.add_paragraph()
    format_paragraph(p, after=4, line=1.0)
    add_text(p, "Integrantes", size=9, bold=True, color=BLUE)
    for line in [
        "Giovana Arenzano da Palma Martins | (11) 99705-9633 | giovana.arenzano@gmail.com",
        "Julianna Rosa Del Cielo | (11) 95197-6037 | juliannajurdc@gmail.com",
        "Karen Mendes Neves de Oliveira | (11) 96494-2847 | karenmendes@wiz.co",
    ]:
        add_body(doc, line, after=3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, before=22, after=0, line=1.0)
    add_text(p, "15 de setembro de 2026", size=9.5, color=GRAY)


def build_report():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    build_cover(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_body_section(body_section)

    # 2 — resumo
    add_page_title(doc, "Relatório técnico", "Resumo executivo", "Visão do problema, da solução entregue e do resultado demonstrado.", page_break=False)
    add_callout(doc, "Proposta.", "O VigiaApp transforma alertas meteorológicos públicos em comunicações preventivas individualizadas para segurados expostos. O MVP percorre o fluxo completo: coleta, interpretação, decisão, geração por IA, revisão e registro do envio.")
    add_heading(doc, "Problema tratado")
    add_body(doc, "Eventos climáticos severos atingem pessoas e patrimônios com pouco tempo de reação. Ao cruzar localização, vigência, produto e coberturas com alertas oficiais, a seguradora pode orientar ações preventivas antes do agravamento do risco e manter uma trilha auditável da comunicação.")
    add_heading(doc, "Resultado do MVP")
    add_metric_cards(doc, [("5", "SEGURADOS FICTÍCIOS", BLUE), ("72 h", "HORIZONTE DE DECISÃO", TEAL), ("2", "PROVEDORES DE IA", NAVY), ("30 min", "CICLO AUTOMÁTICO", BLUE)])
    add_body(doc, "A aplicação funcional integra a API do INMET, valida municípios pelo IBGE, aplica a matriz de riscos e coberturas da Etapa 1, gera mensagens com Gemini 3.5 Flash e utiliza Groq com GPT‑OSS como contingência. A interface permite revisar, regenerar, enviar e consultar o histórico.")
    add_heading(doc, "Escopo demonstrado")
    add_bullet(doc, "Coleta e normalização de avisos meteorológicos reais, sem dados climáticos mockados.")
    add_bullet(doc, "Seleção automática dos segurados elegíveis conforme localização, apólice e cobertura.")
    add_bullet(doc, "Mensagens em português do Brasil adequadas ao canal SMS ou e-mail.")
    add_bullet(doc, "Confirmação de sucesso ou erro no envio e log técnico persistido para falhas.")
    add_callout(doc, "Critério do desafio.", "O edital exige uma demonstração completa e admite simulação de notificação. O MVP registra o resultado do envio no próprio sistema; a conexão com gateway externo de SMS/e-mail fica fora desta versão.", fill="F6F1E8", accent=AMBER)

    # 3 — arquitetura
    add_page_title(doc, "Solução", "Arquitetura do MVP", "Componentes separados por responsabilidade e conectados por contratos simples de dados.")
    add_architecture(doc)
    add_heading(doc, "Camadas")
    add_table(doc, ["Camada", "Responsabilidade", "Implementação"], [
        ("Fontes externas", "Alertas meteorológicos e validação territorial", "API pública do INMET e serviço do IBGE"),
        ("Coleta", "Consultar, validar, normalizar e revisar eventos", "Node.js/TypeScript; histórico de revisões"),
        ("Decisão", "Cruzar evento, município, apólice, produto e cobertura", "Motor determinístico com regra etapa1-v1"),
        ("Geração", "Redigir comunicação personalizada e segura", "Gemini 3.5 Flash; contingência Groq GPT‑OSS"),
        ("Persistência", "Guardar segurados, eventos, comunicações e logs", "SQLite no processo do backend"),
        ("Experiência", "Operar e acompanhar o fluxo", "API Express e painel React/Vite"),
    ], widths=[3.1, 6.5, 7.0], font_size=8.5)
    add_heading(doc, "Princípios de projeto")
    add_bullet(doc, "Separação entre coleta, regras, geração, comunicação e apresentação.")
    add_bullet(doc, "Decisão de elegibilidade determinística; a IA redige, mas não decide cobertura.")
    add_bullet(doc, "Chaves de API somente em variáveis de ambiente; nenhuma credencial no repositório.")
    add_bullet(doc, "Rastreabilidade por IDs, origem do alerta, versão da regra, provedor e modelo usado.")
    add_callout(doc, "Implantação.", "O backend exige um processo persistente para preservar o SQLite e executar o agendador. Hospedagens efêmeras devem usar banco externo e scheduler dedicado.")

    # 4 — agentes
    add_page_title(doc, "Solução", "Agentes e responsabilidades", "Papéis especializados reduzem acoplamento e tornam cada decisão explicável.")
    add_table(doc, ["Agente", "Entrada", "Ação", "Saída"], [
        ("Coletor climático", "Resposta do INMET", "Valida formato, datas, severidade, áreas e revisões", "Evento normalizado"),
        ("Validador territorial", "Município e UF", "Confere referência municipal", "Localidade consistente"),
        ("Motor de decisão", "Evento + segurado + apólice", "Aplica horizonte, vigência e interseção de riscos/coberturas", "Elegível ou descartado, com motivo"),
        ("Redator de IA", "Evento + perfil elegível", "Gera texto preventivo por canal e contexto", "Mensagem revisável"),
        ("Orquestrador", "Ciclo agendado ou comando", "Coordena etapas, reaproveita comunicações e registra erros", "Execução auditável"),
        ("Comunicador", "Mensagem aprovada", "Revalida o alerta e registra sucesso/falha", "Status e histórico"),
    ], widths=[3.0, 3.7, 6.2, 3.7], font_size=8.2)
    add_heading(doc, "Divisão entre regras e IA")
    add_body(doc, "O motor de decisão mantém as condições objetivas fora do modelo generativo. Isso evita que a IA invente elegibilidade ou interprete uma cobertura contratual como certeza de indenização. A IA recebe somente os dados necessários para redigir uma orientação contextualizada.")
    add_heading(doc, "Ciclo operacional")
    add_numbered(doc, "01", "Coleta", "busca alertas ativos no INMET e armazena eventos novos ou revisados.")
    add_numbered(doc, "02", "Decisão", "avalia os cinco segurados fictícios contra cada evento dentro de 72 horas.")
    add_numbered(doc, "03", "Geração", "produz comunicação inédita ou mantém a já existente para evitar duplicidade.")
    add_numbered(doc, "04", "Operação", "expõe resultados no painel e permite revisão, regeneração, envio e histórico.")

    # 5 — fluxo
    add_page_title(doc, "Processamento", "Fluxo ponta a ponta", "Da fonte oficial à mensagem registrada, com guardas em cada transição.")
    steps = [
        ("1", "Consultar fontes", "O backend solicita avisos ativos ao INMET. A carteira e o histórico vêm do SQLite."),
        ("2", "Validar e normalizar", "Campos obrigatórios, datas, municípios, severidade, instruções e riscos são convertidos para o modelo interno."),
        ("3", "Aplicar horizonte", "Somente eventos com início ou vigência dentro da janela operacional de 72 horas seguem para decisão."),
        ("4", "Cruzar segurados", "Município/UF, vigência da apólice, produto e riscos cobertos precisam coincidir com o evento."),
        ("5", "Gerar mensagem", "O Gemini redige por canal; em limite ou indisponibilidade temporária, o Groq assume automaticamente."),
        ("6", "Revisar e enviar", "O usuário revisa ou regenera. Ao enviar, o alerta é revalidado e o sistema registra confirmação ou falha."),
        ("7", "Auditar", "Eventos, revisões, decisão, provedor, modelo, conteúdo, status e logs permanecem consultáveis."),
    ]
    for n, title, text in steps:
        add_numbered(doc, n.zfill(2), title, text)
    add_heading(doc, "Tratamento de repetição e falha")
    add_bullet(doc, "Uma chave de comunicação por evento, revisão e segurado impede duplicações silenciosas.")
    add_bullet(doc, "Revisões do INMET podem gerar nova avaliação quando alteram o risco ou o período.")
    add_bullet(doc, "Falhas externas não produzem conteúdo fictício; a execução informa o erro e mantém o diagnóstico técnico.")

    # 6 — dados
    add_page_title(doc, "Processamento", "Dados externos e persistência", "Origem, uso e evidência guardada pelo sistema.")
    add_table(doc, ["Fonte/conjunto", "Dados usados", "Finalidade", "Tratamento"], [
        ("INMET — avisos ativos", "Tipo, severidade, início/fim, área, riscos e instruções", "Detectar ameaça climática oficial", "Validação de JSON, normalização e revisões"),
        ("IBGE — localidades", "Município, UF e código", "Evitar correspondência territorial ambígua", "Consulta e validação de referência"),
        ("Carteira demonstrativa", "Nome, cidade, produto, coberturas, vigência e canal", "Avaliar impacto individual", "5 perfis fictícios persistidos"),
        ("Modelos de IA", "Contexto mínimo do evento e do segurado", "Redigir mensagem preventiva", "Saída validada e identificada por provedor/modelo"),
        ("SQLite", "Eventos, decisões, mensagens, status e logs", "Histórico e auditoria", "Persistência transacional local"),
    ], widths=[3.2, 5.2, 4.1, 4.2], font_size=8.25)
    add_heading(doc, "Carteira demonstrativa")
    add_body(doc, "A Etapa 1 não fornece uma base de segurados. Para viabilizar o fluxo solicitado, foram criados cinco perfis fictícios, distribuídos entre localidades, produtos, coberturas e canais. Entre eles está Yure Santana, com recebimento por SMS no número informado para a demonstração.")
    add_callout(doc, "Privacidade por desenho.", "O prompt recebe apenas os campos necessários à comunicação. Segredos ficam no ambiente do backend e os logs técnicos não são apresentados como mensagem ao segurado.")
    add_heading(doc, "Registro mínimo por comunicação")
    add_bullet(doc, "evento e revisão de origem; segurado; riscos elegíveis; regra aplicada;")
    add_bullet(doc, "canal; mensagem; provedor e modelo de IA; data de geração;")
    add_bullet(doc, "data e resultado do envio; erro público; detalhe técnico restrito quando houver.")

    # 7 — regras
    add_page_title(doc, "Decisão", "Regras de elegibilidade", "A matriz da Etapa 1 orienta a seleção; o contrato continua sendo a referência de cobertura.")
    add_table(doc, ["Ordem", "Regra", "Condição para avançar"], [
        ("1", "Evento válido", "Alerta normalizado, ativo e com período interpretável"),
        ("2", "Janela temporal", "Evento relevante dentro do horizonte de 72 horas"),
        ("3", "Exposição territorial", "Município e UF do segurado pertencem à área afetada"),
        ("4", "Vigência", "Apólice ativa no período considerado"),
        ("5", "Compatibilidade", "Produto possui relação com ao menos um risco do evento"),
        ("6", "Cobertura declarada", "Há interseção entre riscos elegíveis e coberturas do perfil"),
        ("7", "Não duplicidade", "Ainda não existe comunicação para evento, revisão e segurado"),
    ], widths=[1.6, 4.8, 10.1], font_size=8.55)
    add_heading(doc, "Recorte da matriz Etapa 1")
    add_table(doc, ["Produto", "Riscos climáticos considerados"], [
        ("Agrícola", "seca, calor, chuva, granizo, geada, vendaval, alagamento"),
        ("Residencial", "chuva, alagamento, vendaval, ciclone, granizo, deslizamento, incêndio, raio"),
        ("Empresarial", "chuva, alagamento, vendaval, granizo, incêndio, raio, variação térmica"),
        ("Automóvel", "chuva, alagamento, vendaval, granizo, cheia, deslizamento"),
        ("Engenharia", "chuva, alagamento, vendaval, deslizamento, incêndio"),
        ("Transportes", "chuva, alagamento, vendaval, granizo, cheia, tromba-d’água"),
        ("Paramétrico", "seca, chuva, calor, geada, cheia, variação térmica"),
        ("Responsabilidade civil", "efeitos de chuva, alagamento, vendaval, deslizamento e incêndio quando aplicável"),
    ], widths=[4.0, 12.6], font_size=8.15)
    add_callout(doc, "Interpretação.", "A correspondência sinaliza relevância preventiva. Cobertura efetiva, limites, franquias e exclusões dependem da apólice e da análise de sinistro; o sistema não promete indenização.", fill="F6F1E8", accent=AMBER)

    # 8 — IA
    add_page_title(doc, "Inteligência artificial", "Geração e contingência", "Mensagens úteis, limitadas aos fatos disponíveis e com continuidade entre provedores.")
    add_heading(doc, "Estratégia de modelos")
    add_table(doc, ["Papel", "Provedor/modelo", "Quando é usado"], [
        ("Primário", "Google — Gemini 3.5 Flash", "Primeira tentativa de geração ou regeneração"),
        ("Contingência", "Groq — openai/gpt-oss-120b", "Limite 429, modelo indisponível ou falha temporária do Gemini"),
    ], widths=[3.1, 6.2, 7.3], font_size=8.6)
    add_heading(doc, "Contrato do prompt")
    add_bullet(doc, "Escrever em português do Brasil e adaptar extensão e tom ao canal.")
    add_bullet(doc, "Usar somente o alerta oficial e o perfil fornecido; não inventar datas, locais ou coberturas.")
    add_bullet(doc, "Distinguir previsão/aviso de evento confirmado e informar a fonte INMET.")
    add_bullet(doc, "Fornecer recomendações seguras e acionáveis, além dos números 199 e 193 quando pertinentes.")
    add_bullet(doc, "Evitar promessa de indenização e lembrar que a cobertura depende das condições contratuais.")
    add_heading(doc, "Fallback automático")
    add_numbered(doc, "01", "Tentar Gemini", "envia o contexto estruturado e valida a resposta.")
    add_numbered(doc, "02", "Classificar falha", "limite ou indisponibilidade temporária habilitam a contingência.")
    add_numbered(doc, "03", "Acionar Groq", "repete o mesmo contrato de conteúdo com GPT‑OSS.")
    add_numbered(doc, "04", "Registrar origem", "grava provedor, modelo e indicação de fallback para auditoria.")
    add_callout(doc, "Sem texto de reserva.", "Se os dois provedores falharem, o sistema retorna erro; não substitui a resposta por mensagem mockada ou por um stub.")

    # 9 — experiência
    add_page_title(doc, "Comunicação", "Operação, envio e histórico", "O operador controla a mensagem; o sistema mantém estados coerentes e evidência técnica.")
    add_table(doc, ["Área", "Função"], [
        ("Dashboard", "Resume eventos, segurados, comunicações pendentes e atividade do agente."),
        ("Eventos", "Exibe avisos climáticos coletados, severidade, vigência, municípios e riscos."),
        ("Comunicações", "Filtra destinatários, revisa texto, edita, regenera e executa o envio."),
        ("Histórico", "Mostra mensagens processadas, canal, status, horário e origem da IA."),
    ], widths=[4.0, 12.6], font_size=8.7)
    add_heading(doc, "Semântica do botão Enviar")
    add_body(doc, "Ao confirmar o envio, o backend revalida o alerta na fonte, verifica o endereço de canal, grava a tentativa e responde ao front-end. O usuário recebe uma confirmação quando a operação é aceita. Se houver falha, vê uma mensagem objetiva de que não foi possível enviar, enquanto o detalhe fica em log técnico oculto.")
    add_heading(doc, "Estados apresentados")
    add_table(doc, ["Estado", "Significado"], [
        ("Pendente", "Mensagem gerada e aguardando ação do operador."),
        ("Enviada", "Tentativa aceita e registrada no fluxo do MVP."),
        ("Erro", "Validação ou processamento falhou; ação não foi concluída."),
    ], widths=[4.0, 12.6], font_size=8.7)
    add_callout(doc, "Canal desta versão.", "O desafio permite simular a notificação. Assim, ‘Enviada’ representa o envio processado e registrado pela aplicação. Confirmação de entrega por operadora de SMS/e-mail exige integração futura com um provedor externo.", fill="F6F1E8", accent=AMBER)
    add_heading(doc, "Exemplo de destinatário")
    add_body(doc, "Yure Santana · Belo Horizonte/MG · Seguro residencial · Canal SMS · (71) 99140-8574. O perfil participa do cenário de tempestade quando os riscos do aviso intersectam as coberturas declaradas.")

    # 10 — exemplos
    add_page_title(doc, "Evidências", "Mensagens geradas pelo agente", "Exemplos produzidos durante a validação com alertas reais do INMET.")
    add_heading(doc, "SMS · Yure Santana · Gemini/Groq em contingência")
    add_callout(doc, "Mensagem", "Aviso INMET: Tempestade com risco potencial em Belo Horizonte/MG de 16/09 às 03h até 17/09 às 02h59. Previsão de chuva de 20–30 mm/h, ventos de 40–60 km/h e granizo. Riscos: alagamentos, queda de galhos e danos elétricos. Para sua residência: mantenha janelas e portas fechadas, proteja objetos externos, evite aparelhos ligados à tomada, não se abrigue sob árvores nem estacione perto de torres. Em emergência, procure a Defesa Civil (199) ou os Bombeiros (193).")
    add_body(doc, "Provedor registrado: Groq · Modelo: openai/gpt-oss-120b · Origem do fallback: Gemini.", italic=True, after=10)
    add_heading(doc, "E-mail · seguro empresarial · Gemini")
    add_callout(doc, "Assunto", "Alerta preventivo de tempestade — Belo Horizonte/MG")
    add_body(doc, "Prezado(a) cliente, o INMET emitiu aviso de tempestade com perigo potencial para Belo Horizonte. Há riscos elegíveis de chuva, alagamento e vendaval. Como prevenção, mantenha calhas e ralos desobstruídos, feche portas e janelas externas, evite aparelhos ligados à tomada e não estacione sob árvores ou próximo a torres. Em emergência, ligue 199 ou 193. A cobertura efetiva depende das condições, limites e exclusões da apólice.")
    add_heading(doc, "Qualidade observada")
    add_bullet(doc, "Local, período, severidade e fonte aparecem de forma explícita.")
    add_bullet(doc, "A orientação muda conforme canal, produto e riscos elegíveis.")
    add_bullet(doc, "O texto evita tom alarmista e não promete cobertura ou indenização.")
    add_bullet(doc, "A aplicação registra qual modelo produziu cada mensagem.")

    # 11 — testes
    add_page_title(doc, "Validação", "Duas rodadas de testes", "Verificações automatizadas e exercício do fluxo com serviços reais.")
    add_table(doc, ["Rodada", "Escopo", "Resultado"], [
        ("1", "Backend: regras, coleta, decisão, IA, envio e persistência; front-end: telas, filtros e ações", "Backend 4/4; front-end 76/76; typecheck e build aprovados"),
        ("2", "Reexecução após ajustes de fallback, rótulos e estados; validação do fluxo ponta a ponta", "Mesmos testes aprovados; aplicação funcional com INMET e IA reais"),
    ], widths=[2.2, 9.2, 5.2], font_size=8.5)
    add_heading(doc, "Evidência da execução funcional")
    add_metric_cards(doc, [("8", "EVENTOS COLETADOS", BLUE), ("5", "SEGURADOS AVALIADOS", TEAL), ("4", "ELEGÍVEIS", NAVY), ("0", "ERROS DO CICLO", GREEN)])
    add_body(doc, "Na checagem registrada em 15 de setembro de 2026, o health check informou execução concluída, carteira com cinco segurados e quatro comunicações existentes. O gerador estava disponível via Groq em contingência do Gemini, comprovando a troca automática de provedor.")
    add_heading(doc, "Cenários cobertos")
    add_bullet(doc, "evento válido e segurado elegível; segurado fora da área; apólice fora de vigência;")
    add_bullet(doc, "risco sem cobertura compatível; comunicação já existente; revisão de alerta;")
    add_bullet(doc, "limite ou indisponibilidade do Gemini; geração pelo Groq; falha dos dois provedores;")
    add_bullet(doc, "envio aceito; telefone inválido; falha pública acompanhada de log técnico.")
    add_callout(doc, "Observação.", "Os números de eventos e elegíveis dependem dos avisos ativos do INMET e podem mudar a cada execução.")

    # 12 — aderência
    add_page_title(doc, "Conclusão", "Aderência ao Desafio 5", "Checklist do edital e limites declarados para a primeira versão funcional.")
    add_check(doc, "ATENDE", "API meteorológica pública", "Coleta real de avisos ativos do INMET, com validação e normalização.")
    add_check(doc, "ATENDE", "Identificação automática", "Agendador executa o ciclo e reconhece eventos relevantes e revisões.")
    add_check(doc, "ATENDE", "Regras de decisão", "Município, vigência, horizonte, produto, risco e cobertura são avaliados.")
    add_check(doc, "ATENDE", "Mensagens por IA", "Gemini 3.5 Flash com fallback automático para Groq GPT‑OSS.")
    add_check(doc, "ATENDE", "Envio demonstrado", "Confirmação/erro, status e log técnico são persistidos no MVP.")
    add_check(doc, "ATENDE", "Fluxo completo", "Coleta → decisão → geração → operação → histórico no backend e no painel.")
    add_check(doc, "ATENDE", "Documentação técnica", "Este relatório cobre arquitetura, agentes, tecnologias, fluxo, regras e exemplos.")
    add_heading(doc, "Limitações e próximos passos")
    add_bullet(doc, "Integrar provedor externo para comprovante real de entrega de SMS/e-mail.")
    add_bullet(doc, "Substituir SQLite por banco gerenciado e scheduler dedicado em ambiente serverless.")
    add_bullet(doc, "Adicionar autenticação, perfis de acesso, criptografia operacional e política de retenção.")
    add_bullet(doc, "Validar a matriz com produto jurídico/atuarial e importar carteira autorizada em ambiente seguro.")
    add_bullet(doc, "Formalizar o pacote de submissão com código-fonte, README e licença MIT.")
    add_callout(doc, "Conclusão.", "O VigiaApp atende ao núcleo funcional do Desafio 5: usa dados externos reais, aplica regras claras, produz mensagens personalizadas por IA e demonstra a comunicação ponta a ponta com rastreabilidade. As limitações listadas delimitam a passagem do MVP para produção.")

    # 13 — referências
    add_page_title(doc, "Apêndice", "Referências e acesso", "Documentos de origem, serviços públicos e artefatos do projeto.")
    add_heading(doc, "Documentos de referência")
    add_body(doc, "InsurMinds. Desafios — Desafio 5: Ferramenta Inteligente para Comunicação Proativa com o Segurado. 2026.")
    add_body(doc, "VIL — Visionary Insurance Lab. Etapa 1 — Levantamento de riscos climáticos. 2026.")
    add_heading(doc, "Serviços e documentação técnica")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "INMET — Avisos meteorológicos: "); add_hyperlink(p, "apiprevmet3.inmet.gov.br/avisos/ativos", "https://apiprevmet3.inmet.gov.br/avisos/ativos")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "IBGE — API de localidades: "); add_hyperlink(p, "servicodados.ibge.gov.br", "https://servicodados.ibge.gov.br/api/docs/localidades")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "Google AI Studio: "); add_hyperlink(p, "aistudio.google.com", "https://aistudio.google.com/")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "Groq API: "); add_hyperlink(p, "console.groq.com/docs/api-reference", "https://console.groq.com/docs/api-reference")
    add_heading(doc, "Projeto")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "Repositório público: "); add_hyperlink(p, "github.com/SantanaYure/vigiaApp", "https://github.com/SantanaYure/vigiaApp")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "Aplicação web: "); add_hyperlink(p, "vigia-front-self.vercel.app", "https://vigia-front-self.vercel.app")
    p = doc.add_paragraph(); format_paragraph(p, after=5); add_text(p, "API: "); add_hyperlink(p, "vigia-backend-navy.vercel.app", "https://vigia-backend-navy.vercel.app")
    add_callout(doc, "Versão deste relatório.", "Consolidado em 15 de setembro de 2026 a partir do desafio, da Etapa 1 e da implementação disponível na branch main do repositório.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
